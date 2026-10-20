import os
import json
import re
import uuid
import subprocess
import fitz  # PyMuPDF
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_COLOR_INDEX
from pptx import Presentation
from pptx.util import Pt as PptxPt, Inches as PptxInches
from pptx.dml.color import RGBColor
from docx2pdf import convert

# --- Configuration ---
NOTES_FOLDER = "my_notes"
OUTPUT_FOLDER = "xyxyxy"
TEXT_FILE = os.path.join(NOTES_FOLDER, "test_final.txt")
RESULT_FILE = os.path.join(NOTES_FOLDER, "result_final.txt")

os.makedirs(OUTPUT_FOLDER, exist_ok=True)

# --- Tokenizer Helper ---
def build_token_index(flagged_items):
    """Creates a deduplicated mapping of sensitive strings sorted by length descending."""
    unique_items = {}
    for item in flagged_items:
        raw_text = item.get("text", "").strip()
        if raw_text and raw_text not in unique_items:
            unique_items[raw_text] = item.get("type", "PII")
    
    sorted_texts = sorted(unique_items.keys(), key=len, reverse=True)
    return sorted_texts, unique_items

def tokenize_paragraph(paragraph_text, sorted_texts, type_mapping):
    """Tokenizes text into safe runs identifying PII without altering surrounding words."""
    if not sorted_texts or not paragraph_text.strip():
        return [{"text": paragraph_text, "is_pii": False}]

    replacements = {}
    temp_text = paragraph_text

    # Replace occurrences with unique temporary tokens
    for pii_text in sorted_texts:
        if pii_text in temp_text:
            token = f"<<<{uuid.uuid4().hex}>>>"
            replacements[token] = {"text": pii_text, "type": type_mapping[pii_text]}
            temp_text = temp_text.replace(pii_text, token)

    parts = re.split(r'(<<<[a-f0-9]{32}>>>)', temp_text)
    runs = []
    for part in parts:
        if not part:
            continue
        if part in replacements:
            runs.append({
                "text": replacements[part]["text"],
                "type": replacements[part]["type"],
                "is_pii": True
            })
        else:
            runs.append({"text": part, "is_pii": False})

    return runs

def convert_word_to_pdf(doc_path, pdf_path):
    """Converts DOCX to PDF using MS Word on Windows/macOS or LibreOffice on Linux."""
    try:
        convert(doc_path, pdf_path)
    except Exception as e:
        print("[!] docx2pdf failed, falling back to LibreOffice headless...")
        subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "pdf", doc_path, "--outdir", os.path.dirname(pdf_path)],
            check=True
        )

# --- Document Generators ---
def generate_docx(paragraphs, sorted_texts, type_mapping, out_path, is_safe=False):
    """Generates continuous Word document with 16pt font."""
    doc = Document()

    for para_text in paragraphs:
        if not para_text.strip():
            continue

        runs = tokenize_paragraph(para_text, sorted_texts, type_mapping)
        p = doc.add_paragraph()

        for r in runs:
            if is_safe:
                run_text = f"[{r['type']}]" if r['is_pii'] else r['text']
                run = p.add_run(run_text)
                run.font.size = Pt(16)
            else:
                run = p.add_run(r['text'])
                run.font.size = Pt(16)
                if r['is_pii']:
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    doc.save(out_path)
    print(f"[+] Saved DOCX: {out_path}")

def generate_pptx(paragraphs, sorted_texts, type_mapping, out_path, is_safe=False):
    """
    Generates continuous presentation slides with 13pt font.
    Paginates automatically onto a new slide when word count (~180 words) is reached.
    """
    prs = Presentation()
    words_per_slide_limit = 180

    def create_new_slide():
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank slide
        textbox = slide.shapes.add_textbox(
            PptxInches(0.6), PptxInches(0.6),
            prs.slide_width - PptxInches(1.2), prs.slide_height - PptxInches(1.2)
        )
        tf = textbox.text_frame
        tf.word_wrap = True
        return tf

    current_tf = create_new_slide()
    current_word_count = 0

    for para_text in paragraphs:
        if not para_text.strip():
            continue

        words_in_para = len(para_text.split())

        # If adding this paragraph overflows the slide, advance to next slide
        if current_word_count + words_in_para > words_per_slide_limit and current_word_count > 0:
            current_tf = create_new_slide()
            current_word_count = 0

        p = current_tf.add_paragraph()
        runs = tokenize_paragraph(para_text, sorted_texts, type_mapping)

        for r in runs:
            run = p.add_run()
            if is_safe:
                run.text = f"[{r['type']}]" if r['is_pii'] else r['text']
                run.font.size = PptxPt(13)
            else:
                run.text = r['text']
                run.font.size = PptxPt(13)
                if r['is_pii']:
                    run.font.color.rgb = RGBColor(255, 0, 0)
                    run.font.bold = True

        current_word_count += words_in_para

    prs.save(out_path)
    print(f"[+] Saved PPTX: {out_path}")

def mask_pdf_with_fitz(source_pdf_path, target_pdf_path, sorted_texts, type_mapping):
    """Applies authentic white-box redaction annotations over coordinates using PyMuPDF."""
    doc = fitz.open(source_pdf_path)

    for page in doc:
        for original_text in sorted_texts:
            item_type = type_mapping[original_text]
            rectangles = page.search_for(original_text)
            for rect in rectangles:
                page.add_redact_annot(
                    rect,
                    text=f"[{item_type}]",
                    fill=(1, 1, 1),       # White patch background
                    text_color=(0, 0, 0), # Black text
                    fontsize=10,
                    align=1               # Center align
                )
        page.apply_redactions()

    doc.save(target_pdf_path)
    doc.close()
    print(f"[+] Saved Safe PDF via PyMuPDF Redaction: {target_pdf_path}")

# --- Main Processing Workflow ---
def process_book():
    # 1. Read test_final.txt
    if not os.path.exists(TEXT_FILE):
        print(f"[!] Error: '{TEXT_FILE}' not found.")
        return

    print(f"[*] Loading continuous text from {TEXT_FILE}...")
    with open(TEXT_FILE, "r", encoding="utf-8", errors="ignore") as f:
        raw_content = f.read()

    # Split text into natural paragraphs
    paragraphs = [p.strip() for p in raw_content.splitlines() if p.strip()]
    print(f"[*] Read {len(paragraphs)} paragraph blocks.")

    # 2. Read result_final.txt
    if not os.path.exists(RESULT_FILE):
        print(f"[!] Error: '{RESULT_FILE}' not found.")
        return

    print(f"[*] Loading scan results from {RESULT_FILE}...")
    with open(RESULT_FILE, "r", encoding="utf-8") as f:
        result_json = json.load(f)

    # Handle dictionary or direct list JSON structure
    if isinstance(result_json, dict):
        flagged_items = result_json.get("flagged_items", [])
    elif isinstance(result_json, list):
        flagged_items = result_json
    else:
        flagged_items = []

    print(f"[*] Found {len(flagged_items)} flagged PII item(s).")
    sorted_texts, type_mapping = build_token_index(flagged_items)

    # 3. Generate Original Documents (Highlighted)
    print("\n--- Generating Original Files (Highlighted) ---")
    doc_orig = os.path.join(OUTPUT_FOLDER, "original_final.docx")
    ppt_orig = os.path.join(OUTPUT_FOLDER, "original_final.pptx")
    pdf_orig = os.path.join(OUTPUT_FOLDER, "original_final.pdf")

    generate_docx(paragraphs, sorted_texts, type_mapping, doc_orig, is_safe=False)
    generate_pptx(paragraphs, sorted_texts, type_mapping, ppt_orig, is_safe=False)

    print(f"[*] Converting {doc_orig} to PDF...")
    convert_word_to_pdf(doc_orig, pdf_orig)

    # 4. Generate Safe Documents (Redacted)
    print("\n--- Generating Safe Files (Redacted) ---")
    doc_safe = os.path.join(OUTPUT_FOLDER, "safe_final.docx")
    ppt_safe = os.path.join(OUTPUT_FOLDER, "safe_final.pptx")
    pdf_safe = os.path.join(OUTPUT_FOLDER, "safe_final.pdf")

    generate_docx(paragraphs, sorted_texts, type_mapping, doc_safe, is_safe=True)
    generate_pptx(paragraphs, sorted_texts, type_mapping, ppt_safe, is_safe=True)

    print(f"[*] Redacting original PDF coordinates for {pdf_safe}...")
    mask_pdf_with_fitz(pdf_orig, pdf_safe, sorted_texts, type_mapping)

    print("\n[+] All processing complete! Check the 'xyxyxy' folder for the final files.")

if __name__ == "__main__":
    process_book()