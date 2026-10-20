import os
import json
import re
import uuid
import subprocess
import fitz  # PyMuPDF
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_COLOR_INDEX
from pptx import Presentation
from pptx.util import Pt as PptxPt, Inches as PptxInches
from pptx.dml.color import RGBColor
from docx2pdf import convert

# --- Configuration ---
NOTES_FOLDER = "my_notes"
OUTPUT_FOLDER = "xyxyxy"
PROMPT_FILE = os.path.join(NOTES_FOLDER, "prompt.txt")

os.makedirs(OUTPUT_FOLDER, exist_ok=True)

# --- Tokenizer Helper ---
def tokenize_text(text, flagged_items):
    """
    Safely splits the prompt into a list of dictionaries (runs)
    distinguishing between normal text and PII, handling overlaps.
    """
    if not flagged_items:
        return [{"text": text, "is_pii": False}]
        
    # Deduplicate and sort by length descending (masks longest strings first)
    unique_items = {item['text']: item['type'] for item in flagged_items}
    sorted_texts = sorted(unique_items.keys(), key=len, reverse=True)
    
    replacements = {}
    temp_text = text
    
    # Replace actual PII with unique UUID tokens temporarily
    for pii_text in sorted_texts:
        if pii_text in temp_text:
            token = f"<<<{uuid.uuid4().hex}>>>"
            replacements[token] = {"text": pii_text, "type": unique_items[pii_text]}
            temp_text = temp_text.replace(pii_text, token)
            
    # Split the text by the tokens
    parts = re.split(r'(<<<[a-f0-9]{32}>>>)', temp_text)
    
    runs = []
    for part in parts:
        if not part:
            continue
        if part in replacements:
            runs.append({"text": replacements[part]["text"], "type": replacements[part]["type"], "is_pii": True})
        else:
            runs.append({"text": part, "is_pii": False})
            
    return runs

def convert_word_to_pdf(doc_path, pdf_path):
    """Converts DOCX to PDF (Uses MS Word on Windows/Mac, LibreOffice on Linux)"""
    try:
        convert(doc_path, pdf_path)
    except Exception as e:
        print("[!] docx2pdf failed, trying LibreOffice fallback...")
        subprocess.run(["libreoffice", "--headless", "--convert-to", "pdf", doc_path, "--outdir", os.path.dirname(pdf_path)], check=True)

# --- Core Generators ---
def generate_docx(prompts, all_pii_data, out_path, is_safe=False):
    doc = Document()
    
    for i, (prompt, pii_data) in enumerate(zip(prompts, all_pii_data)):
        runs = tokenize_text(prompt, pii_data)
        
        p = doc.add_paragraph()
        for r in runs:
            if is_safe:
                # Safe Version: Plain text, replaced with [TYPE]
                run_text = f"[{r['type']}]" if r['is_pii'] else r['text']
                run = p.add_run(run_text)
                run.font.size = Pt(16)
            else:
                # Original Version: Keep text, highlight if PII
                run = p.add_run(r['text'])
                run.font.size = Pt(16)
                if r['is_pii']:
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        
        # 2 Prompts per page logic
        if (i + 1) % 2 == 0 and i < len(prompts) - 1:
            # End of pair, insert page break
            doc.add_page_break()
        elif i < len(prompts) - 1:
            # Middle of a pair, add a visual separator
            p_sep = doc.add_paragraph("\n--------------------------------------------------\n")
            p_sep.runs[0].font.size = Pt(16)

    doc.save(out_path)
    print(f"[+] Saved DOCX: {out_path}")

def generate_pptx(prompts, all_pii_data, out_path, is_safe=False):
    prs = Presentation()
    
    for prompt, pii_data in zip(prompts, all_pii_data):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
        
        # Add textbox filling most of the slide
        txBox = slide.shapes.add_textbox(PptxInches(0.5), PptxInches(0.5), prs.slide_width - PptxInches(1), prs.slide_height - PptxInches(1))
        tf = txBox.text_frame
        tf.word_wrap = True
        
        p = tf.add_paragraph()
        runs = tokenize_text(prompt, pii_data)
        
        for r in runs:
            run = p.add_run()
            if is_safe:
                # Safe Version: Plain text, replaced with [TYPE]
                run.text = f"[{r['type']}]" if r['is_pii'] else r['text']
                run.font.size = PptxPt(17)
            else:
                # Original Version: Bold and Red for PII
                run.text = r['text']
                run.font.size = PptxPt(17)
                if r['is_pii']:
                    run.font.color.rgb = RGBColor(255, 0, 0)
                    run.font.bold = True

    prs.save(out_path)
    print(f"[+] Saved PPTX: {out_path}")

def mask_pdf_with_fitz(source_pdf_path, target_pdf_path, all_pii_data):
    """
    Applies authentic PyMuPDF redaction patches over the original PDF.
    This mimics the exact 'sticker' flattening logic from main.py.
    """
    doc = fitz.open(source_pdf_path)
    
    # Flatten all PII items across all prompts and deduplicate
    all_items = []
    for pii_list in all_pii_data:
        all_items.extend(pii_list)
        
    unique_items = {}
    for item in all_items:
        key = item.get("text", "").strip()
        if key and key not in unique_items:
            unique_items[key] = item.get("type", "PII")
            
    # Sort by length descending to redact longer strings first
    sorted_items = sorted(unique_items.items(), key=lambda x: len(x[0]), reverse=True)
    
    for page in doc:
        for original_text, item_type in sorted_items:
            rectangles = page.search_for(original_text)
            for rect in rectangles:
                page.add_redact_annot(
                    rect,
                    text=f"[{item_type}]",
                    fill=(1, 1, 1),       # White sticker background
                    text_color=(0, 0, 0), # Black text
                    fontsize=10,
                    align=1               # Center align
                )
        # Surgically remove original text and flatten the patch
        page.apply_redactions()
        
    doc.save(target_pdf_path)
    doc.close()
    print(f"[+] Saved Safe PDF via PyMuPDF Redaction: {target_pdf_path}")

def process_all():
    # 1. Read the original prompts
    with open(PROMPT_FILE, "r", encoding="utf-8") as f:
        content = f.read()
    
    raw_prompts = re.split(r'\n\s*\n', content)
    prompts = [p.strip() for p in raw_prompts if p.strip()]
    
    print(f"[*] Loaded {len(prompts)} prompts from {PROMPT_FILE}.")
    
    # 2. Read all the JSON result files (p1.txt, p2.txt...)
    all_pii_data = []
    for i in range(1, len(prompts) + 1):
        result_file = os.path.join(NOTES_FOLDER, f"p{i}.txt")
        if os.path.exists(result_file):
            with open(result_file, "r", encoding="utf-8") as f:
                data = json.load(f)
                all_pii_data.append(data.get("flagged_items", []))
        else:
            print(f"[!] Warning: {result_file} missing. Proceeding with no PII for this prompt.")
            all_pii_data.append([])

    # 3. Generate Original Documents (Highlighted)
    print("\n--- Generating Original (Highlighted) Files ---")
    doc_orig = os.path.join(OUTPUT_FOLDER, "original_test.docx")
    ppt_orig = os.path.join(OUTPUT_FOLDER, "original_test.pptx")
    pdf_orig = os.path.join(OUTPUT_FOLDER, "original_test.pdf")
    
    generate_docx(prompts, all_pii_data, doc_orig, is_safe=False)
    generate_pptx(prompts, all_pii_data, ppt_orig, is_safe=False)
    
    print(f"[*] Converting DOCX to PDF for {pdf_orig}...")
    convert_word_to_pdf(doc_orig, pdf_orig)

    # 4. Generate Safe Documents
    print("\n--- Generating Safe (Redacted) Files ---")
    doc_safe = os.path.join(OUTPUT_FOLDER, "safe_test.docx")
    ppt_safe = os.path.join(OUTPUT_FOLDER, "safe_test.pptx")
    pdf_safe = os.path.join(OUTPUT_FOLDER, "safe_test.pdf")
    
    # Generate Safe Word and PPTX normally (plain text substitution)
    generate_docx(prompts, all_pii_data, doc_safe, is_safe=True)
    generate_pptx(prompts, all_pii_data, ppt_safe, is_safe=True)
    
    # CRITICAL CHANGE: Generate Safe PDF by stamping the Authentic original PDF using PyMuPDF!
    print(f"[*] Redacting original PDF coordinates for {pdf_safe}...")
    mask_pdf_with_fitz(pdf_orig, pdf_safe, all_pii_data)
    
    print("\n[+] All document generation complete! Check the 'xyxyxy' folder.")

if __name__ == "__main__":
    process_all()